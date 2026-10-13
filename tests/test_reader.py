import json
import shutil
import sqlite3
import time
from pathlib import Path
import pytest
from PySide6.QtWidgets import QApplication
from PySide6.QtGui import QTextCursor
from PySide6.QtCore import Qt
from PySide6.QtTest import QTest

from perovskite.store import Store, restore
from perovskite.selftest import fixtures
from perovskite.extract import digest
from perovskite.reader import ReaderPanel, load_pdf_page, docx_display, docx_selection, utf16_length, transform_rect
from perovskite.exporting import markdown, excel_export

@pytest.fixture(scope="session")
def qtapp():
    return QApplication.instance() or QApplication([])

@pytest.fixture
def reader_library(tmp_path):
    files=fixtures(tmp_path/"合成 测试文档")
    with Store(tmp_path/"便携 文献库") as store:
        pdf=store.import_file(files["main"])
        docx=store.import_file(files["si_docx"])
        store.associate(docx["file_id"],pdf["article_id"],"si")
        yield store,pdf,docx,files

def wait_reader(panel,app):
    deadline=time.monotonic()+10
    while time.monotonic()<deadline:
        app.processEvents()
        if not panel.active_worker and not panel.pending and not panel.fit_timer.isActive():
            return
        time.sleep(.01)
    raise AssertionError("阅读器渲染未在10秒内结束")

def test_zoom_coordinates_rotation_and_no_original_changes(reader_library,tmp_path):
    store,pdf,docx,files=reader_library
    original=digest(store.path(pdf["file_id"]))
    small=load_pdf_page(store.path(pdf["file_id"]),0,50)
    large=load_pdf_page(store.path(pdf["file_id"]),0,200)
    from PySide6.QtGui import QImage
    assert QImage.fromData(large["png"]).width() >= QImage.fromData(small["png"]).width()*3.9
    assert small["words"][0]["rect"]==large["words"][0]["rect"]
    import fitz
    rotated=tmp_path/"旋转页.pdf"
    with fitz.open(files["main"]) as doc:
        doc[0].set_rotation(90); doc.save(rotated)
    data=load_pdf_page(rotated,0,100)
    word=data["words"][0]
    assert transform_rect(word["rect"],data["rotation"])==pytest.approx(word["display_rect"])
    assert all(0 <= n <= max(data["width"],data["height"]) for n in word["display_rect"])
    assert digest(store.path(pdf["file_id"]))==original
    huge=tmp_path/"大页面.pdf"
    with fitz.open() as doc:
        doc.new_page(width=3000,height=3000); doc.save(huge)
    capped=load_pdf_page(huge,0,400)
    image=QImage.fromData(capped["png"])
    assert capped["limited"] and image.width()*image.height()<=16_010_000

def test_pdf_drag_highlight_zoom_reopen_and_delete(reader_library,qtapp):
    store,pdf,docx,files=reader_library
    panel=ReaderPanel(store); errors=[]; panel.error.connect(errors.append)
    panel.resize(960,900); panel.show(); panel.open_file(pdf["file_id"]); wait_reader(panel,qtapp)
    words=panel.canvas.data["words"]
    start=next(i for i,w in enumerate(words) if w["text"]=="Control")
    end=start+8
    pos1=panel.canvas.screen_rect(words[start]["display_rect"]).center().toPoint()
    pos2=panel.canvas.screen_rect(words[end]["display_rect"]).center().toPoint()
    QTest.mousePress(panel.canvas,Qt.MouseButton.LeftButton,Qt.KeyboardModifier.NoModifier,pos1)
    QTest.mouseMove(panel.canvas,pos2,20)
    QTest.mouseRelease(panel.canvas,Qt.MouseButton.LeftButton,Qt.KeyboardModifier.NoModifier,pos2)
    assert "Control sample annealed" in panel.canvas.selected_payload()["text"]
    assert "10" in panel.canvas.selected_payload()["text"]
    panel.note.setText("准备复现退火条件")
    panel.save_selection()
    assert not errors and len(store.highlights(pdf["file_id"]))==1
    assert store.list_articles("准备复现")
    panel.zoom.setValue(200); wait_reader(panel,qtapp)
    assert panel.zoom.value()==200 and panel.canvas.marks
    assert store.reader_state(pdf["file_id"])["zoom"]==200
    panel.open_file(docx["file_id"]); panel.open_file(pdf["file_id"]); wait_reader(panel,qtapp)
    assert panel.zoom.value()==200 and panel.marks.topLevelItemCount()==1
    panel.marks.setCurrentItem(panel.marks.topLevelItem(0)); panel.delete_mark()
    assert not store.highlights(pdf["file_id"])
    assert not store.list_articles("准备复现")
    assert digest(store.path(pdf["file_id"]))==store.file(pdf["file_id"])["sha256"]
    panel.stop(); wait_reader(panel,qtapp); panel.close()

def test_docx_selected_text_colors_and_unicode_offsets(reader_library,qtapp):
    store,pdf,docx,files=reader_library
    panel=ReaderPanel(store); errors=[]; panel.error.connect(errors.append)
    panel.open_file(docx["file_id"])
    cursor=panel.doc_text.document().find("Control sample annealed at 100 C for 10 min.")
    panel.doc_text.setTextCursor(cursor); panel.color.setCurrentIndex(panel.color.findData("green")); panel.save_selection()
    assert not errors and len(store.highlights(docx["file_id"]))==1
    assert panel.doc_text.extraSelections()
    panel.zoom.setValue(150)
    assert panel.doc_text.font().pointSizeF()>panel.base_font.pointSizeF()
    panel.open_file(docx["file_id"])
    assert len(panel.doc_text.extraSelections())==1
    display,spans=docx_display([{"loc_id":"docx:p0:c0","text":"中文😀句子","section":""}])
    offset=spans[0]["start"]+utf16_length("中文")
    chosen=docx_selection(spans,offset,offset+utf16_length("😀句子"))
    assert chosen=={"text":"😀句子","anchors":[{"loc_id":"docx:p0:c0","start":2,"end":5}]}
    panel.close()

def test_highlights_backup_move_search_and_exports(reader_library,tmp_path):
    store,pdf,docx,files=reader_library
    data=load_pdf_page(store.path(pdf["file_id"]),0)
    word=data["words"][0]
    mark_id=store.save_highlight(pdf["file_id"],"pdf",word["text"],[{"page_index":0,"rect":word["rect"]}],"pink","高亮批注测试")
    store.save_reader_state(pdf["file_id"],0,175,"manual")
    assert "高亮批注测试" in markdown(store,pdf["article_id"])
    excel_export(store,[pdf["article_id"]],tmp_path/"高亮.xlsx")
    from openpyxl import load_workbook
    workbook=load_workbook(tmp_path/"高亮.xlsx")
    assert workbook["高亮批注"].cell(2,10).value=="高亮批注测试"
    archive=tmp_path/"完整备份.zip"; store.backup(archive)
    restore(archive,tmp_path/"恢复库")
    with Store(tmp_path/"恢复库") as restored:
        assert restored.highlights(pdf["file_id"])[0]["mark_id"]==mark_id
        assert restored.reader_state(pdf["file_id"])["zoom"]==175
        assert restored.list_articles("高亮批注测试")
    moved=tmp_path/"移动 数据目录"; shutil.copytree(store.root,moved)
    with Store(moved) as restored:
        assert restored.highlights(pdf["file_id"])[0]["text"]==word["text"]
        assert restored.path(pdf["file_id"]).exists()

def test_migration_from_real_v1_copy_preserves_old_records(reader_library,tmp_path):
    store,pdf,docx,files=reader_library
    old=tmp_path/"旧版本库"; shutil.copytree(store.root,old)
    with sqlite3.connect(old/"library.sqlite") as conn:
        conn.execute("DROP TABLE highlights"); conn.execute("DROP TABLE reader_state"); conn.execute("PRAGMA user_version=1")
    with Store(old) as upgraded:
        assert upgraded.db.execute("PRAGMA user_version").fetchone()[0]==2
        assert upgraded.article(pdf["article_id"])
        assert upgraded.file(pdf["file_id"])["sha256"]==store.file(pdf["file_id"])["sha256"]
        assert list((old/"migrations").glob("library-v1-*.sqlite"))
        backup=list((old/"migrations").glob("*.sqlite"))[0]
        with sqlite3.connect(backup) as conn: assert conn.execute("PRAGMA user_version").fetchone()[0]==1

@pytest.mark.parametrize("anchor,color", [({"page_index":99,"rect":[0,0,10,10]},"yellow"),({"page_index":0,"rect":[0,0,float('nan'),10]},"yellow"),({"page_index":0,"rect":[5,5,0,0]},"yellow"),({"page_index":0,"rect":[0,0,10,10]},"red")])
def test_bad_highlight_positions_and_colors_rejected(reader_library,anchor,color):
    store,pdf,docx,files=reader_library
    with pytest.raises(ValueError): store.save_highlight(pdf["file_id"],"pdf","test",[anchor],color)
    assert not store.highlights(pdf["file_id"])

def test_latest_render_wins_and_scan_page_status(reader_library,qtapp):
    store,pdf,docx,files=reader_library
    scanned=store.import_file(files["scanned"])
    panel=ReaderPanel(store); panel.resize(960,900); panel.show(); panel.open_file(pdf["file_id"])
    panel.zoom.setValue(200); panel.zoom.setValue(50); panel.zoom.setValue(150)
    wait_reader(panel,qtapp)
    assert panel.zoom.value()==150
    panel.query.setText("Control sample annealed"); panel.find_sentence(); wait_reader(panel,qtapp)
    assert panel.canvas.selected_payload()["anchors"]
    panel.open_file(scanned["file_id"]); wait_reader(panel,qtapp)
    assert not panel.canvas.data["words"] and "无可选文字" in panel.status.text()
    panel.stop(); panel.close()
