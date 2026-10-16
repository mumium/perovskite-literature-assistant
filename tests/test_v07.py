"""Synthetic manual-classification and whole-document reader acceptance."""
import shutil
import sqlite3
import time
from contextlib import closing
import pytest
from PySide6.QtCore import Qt,QPointF
from PySide6.QtGui import QDropEvent
from PySide6.QtWidgets import QApplication
from perovskite.store import Store,restore,UNFILED,DB_VERSION
from perovskite.selftest import fixtures
from perovskite.categories import article_mime,read_article_mime
from perovskite.gui import Window
from perovskite.reader import ReaderPanel
from perovskite.exporting import excel_export,markdown

@pytest.fixture(scope='module')
def qt(): return QApplication.instance() or QApplication([])

def wait(panel,qt):
    deadline=time.monotonic()+20
    while time.monotonic()<deadline:
        qt.processEvents()
        if not panel.active_worker and not panel.pending and not panel.fit_timer.isActive(): return
        time.sleep(.005)
    raise AssertionError('全文预加载未结束')

def pdf(path,count):
    import fitz
    with fitz.open() as doc:
        for i in range(count): doc.new_page().insert_text((40,60),f'SYNTHETIC TEST ONLY page {i+1} unique information')
        doc.save(path)
    return path

def test_categories_crud_multi_membership_atomic_validation_and_exports(tmp_path):
    source=fixtures(tmp_path/'合成 分类文档')
    with Store(tmp_path/'中文 数据目录') as store:
        a=store.import_file(source['main']); b=store.import_file(source['si_pdf'])
        ids=[a['article_id'],b['article_id']]; hashes=[store.file(x['file_id'])['sha256'] for x in [a,b]]
        c=store.create_category('手动分类 中文'); d=store.create_category('复现实验')
        assert store.assign_category(ids+[ids[0]],c)==2
        assert store.assign_category(ids,c)==0
        assert store.assign_category([ids[0]],d)==1
        assert len(store.article_categories(ids[0]))==2
        with pytest.raises(ValueError): store.assign_category([ids[0],'错误文献ID'],d)
        assert len(store.list_articles(category_id=c))==2
        assert store.list_articles(category_id=UNFILED)==[]
        assert len(store.list_articles('synthetic',category_id=d))==1
        with pytest.raises(ValueError): store.create_category(' 手动分类 中文 ')
        with pytest.raises(ValueError): store.rename_category(d,'')
        store.rename_category(c,'准备重读'); store.remove_category([ids[1]],c)
        assert len(store.list_articles(category_id=UNFILED))==1
        target=tmp_path/'类别.xlsx'; excel_export(store,ids,target)
        from openpyxl import load_workbook
        book=load_workbook(target); assert book['手动分类'].max_row==3; book.close()
        assert '准备重读；复现实验' in markdown(store,ids[0])
        store.delete_category(c); assert len(store.article_categories(ids[0]))==1
        assert [store.file(x['file_id'])['sha256'] for x in [a,b]]==hashes
        assert all(store.path(x['file_id']).exists() for x in [a,b])
        assert store.rows("SELECT * FROM history WHERE scope='categories'")

def test_category_upgrade_v2_backup_portable_restore(tmp_path):
    source=fixtures(tmp_path/'迁移合成')
    root=tmp_path/'旧库'
    with Store(root) as store:
        a=store.import_file(source['main']); aid=a['article_id']
    with closing(sqlite3.connect(root/'library.sqlite')) as db,db:
        db.execute('DROP TABLE article_categories'); db.execute('DROP TABLE categories'); db.execute('PRAGMA user_version=2')
    with Store(root) as store:
        assert store.db.execute('PRAGMA user_version').fetchone()[0]==DB_VERSION
        snapshot=next((root/'migrations').glob('library-v2-*.sqlite'))
        with closing(sqlite3.connect(snapshot)) as db: assert db.execute('PRAGMA user_version').fetchone()[0]==2
        cid=store.create_category('移动后仍保留'); store.assign_category([aid],cid)
        store.backup(tmp_path/'备份.zip')
    moved=tmp_path/'移动 中文 空格'; shutil.move(str(root),str(moved))
    with Store(moved) as store:
        assert store.list_articles(category_id=cid)[0]['article_id']==aid
        assert store.path(a['file_id']).exists()
    restore(tmp_path/'备份.zip',tmp_path/'恢复分类库')
    with Store(tmp_path/'恢复分类库') as store: assert store.article_categories(aid)[0]['name']=='移动后仍保留'

def test_actual_table_multiselect_drop_category_filter_and_tab(tmp_path,qt):
    source=fixtures(tmp_path/'拖入分类合成')
    root=tmp_path/'拖放库'
    with Store(root) as store:
        a=store.import_file(source['main']); b=store.import_file(source['si_pdf']); cid=store.create_category('收藏 中文')
    window=Window(root); window.resize(1280,720); window.show(); wait(window.reader,qt)
    window.tabs.setCurrentWidget(window.meta_tab); window.table.selectAll()
    selected=window.selected_ids(); assert len(selected)==2
    item=next(window.category_tree.topLevelItem(i) for i in range(window.category_tree.topLevelItemCount()) if window.category_tree.topLevelItem(i).data(0,Qt.ItemDataRole.UserRole)==cid)
    mime=window.table.selection_mime(); assert read_article_mime(mime,root.resolve())==selected
    assert read_article_mime(mime,tmp_path/'另一个库') is None
    event=QDropEvent(QPointF(window.category_tree.visualItemRect(item).center()),Qt.DropAction.CopyAction,mime,Qt.MouseButton.LeftButton,Qt.KeyboardModifier.NoModifier)
    window.category_tree.dropEvent(event); assert event.isAccepted()
    assert len(window.store.list_articles(category_id=cid))==2 and window.tabs.currentWidget() is window.meta_tab
    item=next(window.category_tree.topLevelItem(i) for i in range(window.category_tree.topLevelItemCount()) if window.category_tree.topLevelItem(i).data(0,Qt.ItemDataRole.UserRole)==cid)
    window.category_clicked(item); assert window.table.rowCount()==2 and window.category_id==cid
    window.remove_from_category(); assert window.table.rowCount()==1
    window.clear_filters(); assert window.table.rowCount()==2
    window.reader.stop(); wait(window.reader,qt); window.close(); qt.processEvents()

def test_full_preload_budget_retained_images_no_jump_or_zoom_render(tmp_path,qt):
    source=pdf(tmp_path/'SYNTHETIC 35 pages.pdf',35)
    with Store(tmp_path/'全文库') as store:
        a=store.import_file(source); panel=ReaderPanel(store); panel.resize(900,750); panel.show()
        panel.pdf_view.cache_limit=8*1024**2
        states=[]; panel.pdf_view.loadingChanged.connect(lambda done,total,state:states.append((done,total,state)))
        panel.open_file(a['file_id']); wait(panel,qt); view=panel.pdf_view
        assert view.full_ready and view.render_count==35 and len(view.cache)==35
        assert view.info['reduced'] and view.cache_bytes<=view.cache_limit
        assert all(p.data and not p._image.isNull() for p in view.pages)
        assert states[-1]==(35,35,'ready') and panel.preload_panel.isHidden()
        before=view.render_count
        for i in [34,0,20,2,33,0]: view.navigate(i); qt.processEvents()
        panel.zoom.setValue(400); wait(panel,qt)
        assert view.render_count==before and all(not p._image.isNull() for p in view.pages)
        view.clear(); wait(panel,qt); assert not view.cache and view.cache_bytes==0
        panel.close()

def test_preload_pause_resume_retains_completed_and_retries_failures(tmp_path,qt,monkeypatch):
    import perovskite.reader as reader
    original=reader.load_pdf_page; failed=set()
    def controlled(path,index,*args,**kwargs):
        time.sleep(.015)
        if index==6 and index not in failed: failed.add(index); raise ValueError('明确模拟单页渲染失败')
        return original(path,index,*args,**kwargs)
    monkeypatch.setattr(reader,'load_pdf_page',controlled)
    source=pdf(tmp_path/'SYNTHETIC pause 16 pages.pdf',16)
    with Store(tmp_path/'暂停库') as store:
        a=store.import_file(source); panel=ReaderPanel(store); panel.show(); view=panel.pdf_view
        def pause(done,total,state):
            if done==3 and state=='loading': view.pause_loading()
        view.loadingChanged.connect(pause); panel.open_file(a['file_id']); wait(panel,qt)
        assert view.paused and not view.full_ready and 3<=len(view.cache)<16
        complete={k[1]:v[0]['image'].cacheKey() for k,v in view.cache.items()}
        view.loadingChanged.disconnect(pause); panel.toggle_preload(); wait(panel,qt)
        assert view.failed_pages=={6} and not view.full_ready and panel.preload_button.text()=='重试预载'
        panel.toggle_preload(); wait(panel,qt)
        assert view.full_ready and view.render_count==16 and not view.failed_pages
        assert all(view.cached(i,100)['image'].cacheKey()==key for i,key in complete.items())
        panel.stop(); wait(panel,qt); panel.close()

def test_rapid_document_switch_cannot_mix_preloaded_pages(tmp_path,qt):
    with Store(tmp_path/'切换库') as store:
        a=store.import_file(pdf(tmp_path/'SYNTHETIC many.pdf',25)); b=store.import_file(pdf(tmp_path/'SYNTHETIC other.pdf',2))
        panel=ReaderPanel(store); panel.show(); panel.open_file(a['file_id']); qt.processEvents(); panel.open_file(b['file_id']); wait(panel,qt)
        view=panel.pdf_view
        assert view.full_ready and len(view.pages)==2 and len(view.cache)==2
        assert all(k[0]==b['file_id'] for k in view.cache)
        panel.stop(); wait(panel,qt); panel.close()
