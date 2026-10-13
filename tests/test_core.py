import copy
import json
import shutil
import socket
import zipfile
from pathlib import Path
import pytest

from perovskite.store import Store, restore, validate_metadata, dumps
from perovskite.extract import extract, candidates, split_location, digest
from perovskite.exchange import export_package, preview, apply_result, strict_json, import_results
from perovskite.selftest import fixtures, result_fixture
from perovskite.exporting import excel_export, markdown

@pytest.fixture
def library(tmp_path, monkeypatch):
    def offline(*a, **kw):
        raise AssertionError("离线流程不应联网")
    monkeypatch.setattr(socket, "create_connection", offline)
    source = fixtures(tmp_path / "合成 测试文档")
    with Store(tmp_path / "文献 数据") as store:
        imported = {k: store.import_file(p) for k,p in source.items()}
        aid = imported["main"]["article_id"]
        for k, role in [("main", "main"), ("si_pdf", "si"), ("si_docx", "si")]:
            store.associate(imported[k]["file_id"], aid, role)
        manifest = export_package(store, aid, tmp_path / "分析包.zip")
        result = result_fixture(store, manifest)
        yield store, source, imported, aid, result

def test_import_states_hash_and_versions(library, tmp_path):
    store, source, imported, aid, value = library
    assert store.import_file(source["main"])["status"] == "duplicate"
    assert len(store.files(aid)) == 3
    assert [imported[k]["status"] for k in ["old", "scanned", "corrupt", "encrypted"]] == ["unsupported", "no_text", "error", "error"]
    assert store.article(aid)["metadata"]["doi"] == "10.9999/synthetic-test"
    modified = tmp_path / "另一版本.pdf"
    import fitz
    with fitz.open(source["main"]) as doc:
        doc[0].insert_text((50, 300), "Different version")
        doc.save(modified)
    new = store.import_file(modified)
    assert new["status"] == "ready" and new["pairing_candidates"] >= 1
    assert new["file_id"] != imported["main"]["file_id"]
    assert new["article_id"] != aid
    assert digest(source["main"]) == store.file(imported["main"]["file_id"])["sha256"]

def test_references_doi_is_ignored():
    assert candidates([{"page_index":0,"text":"Current paper title\nReferences\nDOI: 10.1234/wrong"}], "paper.pdf")["doi"] is None
    assert candidates([{"page_index":0,"text":"Current title\n10.1234/unlabelled"}], "paper.pdf")["doi"] is None
    assert candidates([{"page_index":0,"text":"DOI: 10.1234/a\nDOI: 10.1234/b"}], "paper.pdf")["doi"] is None
    assert candidates([{"page_index":0,"text":"Current paper title\nReferences [1]\nDOI: 10.1234/wrong"}], "paper.pdf")["doi"] is None
    assert candidates([{"page_index":0,"text":"Some author, Journal 2024, https://doi.org/10.1234/citation"}], "paper.pdf")["doi"] is None

def test_pair_correction_and_move(library, tmp_path):
    store, source, imported, aid, value = library
    old_id = imported["si_docx"]["file_id"]
    new_id = store.detach(old_id)
    assert len(store.files(aid)) == 2 and len(store.files(new_id)) == 1
    with pytest.raises(ValueError, match="清单|配对"):
        preview(store, value)
    store.associate(old_id, aid, "si")
    apply_result(store, value)
    target = tmp_path / "移动 文献库"
    store.db.commit()
    shutil.copytree(store.root, target)
    with Store(target) as moved:
        assert moved.path(old_id).exists()
        assert moved.section(aid, "story")
        assert any(r["article_id"] == aid for r in moved.list_articles("5%"))

def test_roundtrip_protection_partial_and_history(library):
    store, source, imported, aid, value = library
    assert all(c["matched"] for c in preview(store, value)["checks"])
    assert apply_result(store, value)["status"] == "applied"
    assert apply_result(store, value)["status"] == "duplicate"
    section = store.section(aid, "story")
    with store.db:
        store.save_section(aid, "story", section["content"], section["evidence"], section["checks"], "人工修订", True)
    new = copy.deepcopy(value)
    new["sections"]["story"]["one_sentence"][0]["text"] = "新分析结果不得覆盖核验"
    result = apply_result(store, new)
    assert result["protected"] == ["story"] and result["applied"] == ["process"]
    assert store.section(aid, "story")["content"] == section["content"]
    assert store.rows("SELECT * FROM history WHERE article_id=?", (aid,))
    process = store.section(aid, "process")["content"]
    assert process["samples"][0]["steps"][0]["parameters"][1]["original_value"] is None
    params = process["samples"][1]["steps"][0]["parameters"]
    assert params[0]["conflict_group"] == params[1]["conflict_group"]
    assert params[-1]["basis"] is None

@pytest.mark.parametrize("mutation", ["article", "hash", "package", "foreign_file", "location", "type", "sample", "nan", "conversion", "evidence", "version", "docx_page"])
def test_invalid_results_rejected(library, mutation):
    store, source, imported, aid, value = library
    bad = copy.deepcopy(value)
    if mutation == "article": bad["article_id"] = imported["old"]["article_id"]
    if mutation == "hash": bad["file_hashes"][0]["sha256"] = "0"*64
    if mutation == "package": bad["package_id"] = "unknown"
    if mutation == "foreign_file": bad["evidence"][0]["file_id"] = imported["old"]["file_id"]
    if mutation == "location": bad["evidence"][0]["loc_id"] = "pdf:p999:c0"
    if mutation == "type": bad["sections"]["story"]["one_sentence"] = "invalid"
    if mutation == "sample": bad["sections"]["process"]["samples"][1]["sample_id"] = "control"
    if mutation == "nan":
        with pytest.raises(ValueError): strict_json('{"x":NaN}')
        return
    if mutation == "conversion": bad["sections"]["process"]["samples"][0]["steps"][0]["parameters"][0]["normalized_value"] = 373.15
    if mutation == "evidence": bad["sections"]["story"]["one_sentence"][0]["evidence_ids"] = ["undefined"]
    if mutation == "version": bad["schema_version"] = "99.0"
    if mutation == "docx_page": next(e for e in bad["evidence"] if e["loc_id"].startswith("docx:"))["printed_page"] = "1"
    with pytest.raises(ValueError): preview(store, bad)
    assert store.section(aid, "story") is None

def test_unmatched_quotes_not_verified(library):
    store, source, imported, aid, value = library
    value["evidence"][0]["quote"] = "This quote was hallucinated."
    report = preview(store, value)
    assert not report["checks"][0]["matched"]
    assert report["warnings"]
    apply_result(store, value)
    assert store.section(aid, "story")["human_verified"] == 0
    assert not store.section(aid, "story")["checks"][0]["matched"]
    assert any(r["article_id"]==aid for r in store.list_articles(filters={"state":"引用待核验"}))

def test_selected_sections_can_be_imported_later(library):
    store, source, imported, aid, value = library
    apply_result(store, value, ["story"])
    assert store.section(aid, "process") is None
    assert "process" in apply_result(store, value, ["process"])["applied"]
    assert store.section(aid, "process")

def test_batch_errors_and_exports_backup(library, tmp_path):
    store, source, imported, aid, value = library
    good = tmp_path / "good.json"; good.write_text(dumps(value), encoding="utf-8")
    bad = tmp_path / "bad.json"; bad.write_text("not json", encoding="utf-8")
    reports = import_results(store, [bad, good])
    assert [r["status"] for r in reports] == ["error", "applied"]
    target = tmp_path / "导出.xlsx"
    excel_export(store, [aid], target)
    from openpyxl import load_workbook
    book = load_workbook(target)
    assert len(book.sheetnames) == 8
    assert book["样品工艺"].max_row == 6
    assert book["证据"].max_row > 3
    assert "正文和SI" in markdown(store, aid)
    backup = tmp_path / "backup.zip"; store.backup(backup)
    restored = tmp_path / "恢复库"; restore(backup, restored)
    with Store(restored) as new:
        assert len(new.files(aid)) == 3
        assert new.section(aid, "process")
    with pytest.raises(ValueError): restore(backup, restored)

def test_hash_tamper_backup_traversal_and_cancel(library, tmp_path):
    store, source, imported, aid, value = library
    path = store.path(imported["main"]["file_id"])
    path.write_bytes(path.read_bytes()+b"tampered")
    with pytest.raises(ValueError, match="哈希"):
        preview(store, value)
    with pytest.raises(InterruptedError):
        store.import_file(source["main"], lambda: True)
    archive = tmp_path / "unsafe.zip"
    with zipfile.ZipFile(archive, "w") as z:
        z.writestr("library.sqlite", b"invalid"); z.writestr("../escape.txt", "unsafe")
    with pytest.raises(ValueError, match="不安全"):
        restore(archive, tmp_path / "safe")
    assert not (tmp_path / "escape.txt").exists()

def test_full_chunks_docx_and_metadata(library):
    store, source, imported, aid, value = library
    text = "x"*25001
    locations = split_location("pdf:p0", text)
    assert "".join(x["text"] for x in locations) == text
    docx = extract(source["si_docx"])
    assert any(l["kind"] == "table_row" and "5%" in l["text"] for l in docx["locations"])
    assert not any("page_index" in l for l in docx["locations"])
    assert validate_metadata({"online_date":"2024", "print_date":"2025-03"})
    with pytest.raises(ValueError): validate_metadata({"online_date":"2024-02-30"})
    with pytest.raises(ValueError): strict_json('{"x":1,"x":2}')
    meta = store.article(aid)["metadata"]
    meta["year"] = "2024"; meta["tags"] = ["缺陷钝化"]
    store.save_metadata(aid, meta, "中文 人工笔记")
    assert store.list_articles(filters={"year":"2024", "tags":"缺陷钝化"})
    assert any(a["article_id"] == aid for a in store.list_articles("中文"))

def test_all_pdf_pages_partial_and_long_title(tmp_path):
    import fitz
    path = tmp_path / "中文 路径 多页.pdf"
    with fitz.open() as doc:
        doc.new_page().insert_text((40,50), "SYNTHETIC TEST title\nDOI: 10.9999/multipage")
        doc.new_page()
        doc.new_page().insert_text((40,50), "LAST PAGE unique process evidence - not an abstract")
        doc.save(path)
    result = extract(path)
    assert result["status"] == "partial"
    assert any(l["page_index"]==2 and "LAST PAGE" in l["text"] for l in result["locations"])
    with Store(tmp_path / "data") as store:
        imported = store.import_file(path)
        aid = imported["article_id"]
        metadata = store.article(aid)["metadata"]
        metadata["title"] = "这是长标题 " * 250
        store.save_metadata(aid, metadata, "")
        assert store.article(aid)["metadata"]["title"] == metadata["title"]
        assert store.list_articles("LAST PAGE")
