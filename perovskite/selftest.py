"""Synthetic fixtures only: these are NOT real papers or scientific claims."""
import copy
import json
import shutil
from pathlib import Path
from .store import Store, atomic_json
from .exchange import export_package, apply_result, preview
from .exporting import excel_export, markdown
from .schema import STORY_FIELDS, template

def fixtures(folder):
    import fitz
    from docx import Document
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    def pdf(name, lines):
        path = folder / name
        with fitz.open() as doc:
            page = doc.new_page()
            page.insert_text((50, 60), "\n".join(lines), fontsize=11)
            doc.save(path)
        return path
    main = pdf("测试 正文 synthetic.pdf", ["SYNTHETIC TEST DOCUMENT - NOT A REAL PAPER", "Test perovskite sample processing", "DOI: 10.9999/synthetic-test", "Authors: Test Alpha, Test Beta", "Control sample annealed at 100 C for 10 min.", "Optimized sample annealed at 110 C for 10 min.", "Test PCE was 20.0 percent under synthetic conditions.", "References", "DOI: 10.9999/reference-should-not-be-used"])
    si_pdf = pdf("测试 SI synthetic.pdf", ["Supporting Information - SYNTHETIC TEST ONLY", "Test perovskite sample processing", "DOI: 10.9999/synthetic-test", "Optimized sample annealed at 120 C for 10 min.", "Humidity was not reported."])
    si_docx = folder / "测试 SI 表格 synthetic.docx"
    doc = Document()
    doc.add_paragraph("SYNTHETIC TEST ONLY", style="Title")
    doc.add_paragraph("Supporting Information")
    doc.add_paragraph("DOI: 10.9999/synthetic-test")
    doc.add_heading("Test procedures", 1)
    doc.add_paragraph("Control sample annealed at 100 C for 10 min.")
    table = doc.add_table(rows=1, cols=3)
    for cell, value in zip(table.rows[0].cells, ["Sample", "Additive", "Basis"]):
        cell.text = value
    row = table.add_row()
    for cell, value in zip(row.cells, ["Optimized", "5%", "Not reported"]):
        cell.text = value
    doc.save(si_docx)
    old = folder / "测试 legacy.doc"
    old.write_bytes(b"SYNTHETIC unsupported legacy document")
    scanned = folder / "测试 无文字.pdf"
    with fitz.open() as doc:
        doc.new_page(); doc.save(scanned)
    corrupt = folder / "测试 损坏.pdf"
    corrupt.write_bytes(b"THIS IS NOT PDF")
    encrypted = folder / "测试 加密.pdf"
    with fitz.open(main) as doc:
        doc.save(encrypted, encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw="owner", user_pw="password")
    return {"main": main, "si_pdf": si_pdf, "si_docx": si_docx, "old": old, "scanned": scanned, "corrupt": corrupt, "encrypted": encrypted}

def result_fixture(store, manifest):
    result = template(manifest["package_id"], manifest["article_id"], manifest["files"])
    main = next(f for f in manifest["files"] if f["role"] == "main")
    si = next(f for f in manifest["files"] if f["name"].endswith(".pdf") and f["role"] == "si")
    result["evidence"] = [
        {"evidence_id": "E1", "file_id": main["file_id"], "loc_id": "pdf:p0:c0", "quote": "Control sample annealed at 100 C for 10 min.", "figure": None, "printed_page": None},
        {"evidence_id": "E2", "file_id": main["file_id"], "loc_id": "pdf:p0:c0", "quote": "Optimized sample annealed at 110 C for 10 min.", "figure": None, "printed_page": None},
        {"evidence_id": "E3", "file_id": si["file_id"], "loc_id": "pdf:p0:c0", "quote": "Optimized sample annealed at 120 C for 10 min.", "figure": None, "printed_page": None},
    ]
    story = result["sections"]["story"]
    story["one_sentence"] = [{"text": "合成测试文档比较两个虚构样品的退火工艺。", "kind": "author_claim", "evidence_ids": ["E1", "E2"]}]
    story["limitations"] = [{"text": "正文和SI优化组温度冲突；全部数据为测试数据。", "kind": "direct_result", "evidence_ids": ["E2", "E3"]}]
    story["relevance"] = [{"text": "此测试结果不能用于实际实验建议。", "kind": "ai_inference", "evidence_ids": []}]
    def parameter(name, value, unit, evidence, conflict=None):
        return {"name": name, "original_value": value, "original_unit": unit, "normalized_value": None, "normalized_unit": None, "basis": None, "conversion_note": None, "conflict_group": conflict, "evidence_ids": evidence}
    result["sections"]["process"] = {"samples": [
        {"sample_id": "control", "label": "Control", "group": "control", "steps": [{"order": 1, "category": "退火", "description": "合成测试", "parameters": [parameter("温度", 100, "C", ["E1"]), parameter("湿度", None, None, [])], "evidence_ids": ["E1"]}]},
        {"sample_id": "optimized", "label": "Optimized", "group": "optimized", "steps": [{"order": 1, "category": "退火", "description": "正文与SI冲突，两值均保留", "parameters": [parameter("温度（正文）", 110, "C", ["E2"], "anneal-temperature"), parameter("温度（SI）", 120, "C", ["E3"], "anneal-temperature"), parameter("添加剂比例", "5%", None, ["E3"])], "evidence_ids": ["E2", "E3"]}]},
    ], "performance": [{"record_id": "synthetic-pce", "sample_id": "optimized", "measurement_kind": "合成测试PCE（非真实结果）", "conditions": {"synthetic": True, "area": None, "scan_direction": None}, "values": [parameter("PCE", 20.0, "%", ["E5"]), parameter("统计样本数", None, None, [])], "evidence_ids": ["E5"]}]}
    result["evidence"].append({"evidence_id": "E5", "file_id": main["file_id"], "loc_id": "pdf:p0:c0", "quote": "Test PCE was 20.0 percent under synthetic conditions.", "figure": None, "printed_page": None})
    # Do not falsely attribute additive evidence: use actual DOCX table row.
    docx = next(f for f in manifest["files"] if f["name"].endswith(".docx"))
    loc = next(l for l in docx["locations"] if l["kind"] == "table_row" and "5%" in l["text"])
    result["evidence"].append({"evidence_id": "E4", "file_id": docx["file_id"], "loc_id": loc["loc_id"], "quote": loc["text"], "figure": None, "printed_page": None})
    result["sections"]["process"]["samples"][1]["steps"][0]["parameters"][-1]["evidence_ids"] = ["E4"]
    return result

def run_selftest(root):
    """Invokable in the frozen exe; requires no Python/Office/network at runtime."""
    root = Path(root).resolve()
    root.mkdir(parents=True, exist_ok=True)
    if any(root.iterdir()):
        raise ValueError("自检请指定新的空输出目录，避免覆盖已存在的文件")
    source = root / "明确标注的测试文档"
    files = fixtures(source)
    data = root / "数据 中文 空格"
    checks = []
    with Store(data) as store:
        imported = {k: store.import_file(p) for k, p in files.items()}
        aid = imported["main"]["article_id"]
        assert store.import_file(files["main"])["status"] == "duplicate"
        checks.append("SHA256重复跳过")
        for key, role in [("main", "main"), ("si_pdf", "si"), ("si_docx", "si")]:
            store.associate(imported[key]["file_id"], aid, role)
        assert len(store.files(aid)) == 3
        assert imported["old"]["status"] == "unsupported"
        assert imported["scanned"]["status"] == "no_text"
        assert imported["corrupt"]["status"] == "error"
        assert imported["encrypted"]["status"] == "error"
        checks.append("PDF+DOCX+多SI关联、DOC/扫描/损坏/加密状态")
        manifest = export_package(store, aid, root / "测试分析包.zip")
        result = result_fixture(store, manifest)
        atomic_json(root / "测试返回结果.json", result)
        assert all(c["matched"] for c in preview(store, result)["checks"])
        assert apply_result(store, result)["status"] == "applied"
        assert apply_result(store, result)["status"] == "duplicate"
        checks.append("分析包交换、引用匹配、重复结果识别")
        assert any(a["article_id"] == aid for a in store.list_articles("合成测试"))
        assert any(a["article_id"] == aid for a in store.list_articles("5%"))
        checks.append("中文阅读卡与SI全文检索")
        import fitz
        with fitz.open(store.path(imported["main"]["file_id"])) as doc:
            doc[0].get_pixmap().save(root / "PDF预览.png")
        excel_export(store, [aid], root / "测试关联表.xlsx")
        (root / "测试阅读卡.md").write_text(markdown(store, aid), encoding="utf-8")
        store.backup(root / "测试完整备份.zip")
        checks.append("PDF渲染、Excel/Markdown与完整备份")
    moved = root / "移动后的数据 目录"
    shutil.copytree(data, moved, dirs_exist_ok=True)
    with Store(moved) as store:
        assert store.path(imported["main"]["file_id"]).exists()
        assert store.section(aid, "process")["content"]["samples"][0]["label"] == "Control"
    checks.append("中文路径、数据搬移、重启持久化")
    report = {"passed": True, "synthetic_only": True, "checks": checks, "article_id": aid, "data_dir": str(data)}
    atomic_json(root / "selftest-report.json", report)
    return report
