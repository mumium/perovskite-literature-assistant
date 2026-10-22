import json
from pathlib import Path
from .schema import STORY_FIELDS
from .store import dumps

def markdown(store, article_id):
    a = store.article(article_id)
    text = ["# " + (a["metadata"].get("title_zh") or a["metadata"].get("title") or article_id),
            f"article_id: {article_id}", f"DOI: {a['metadata'].get('doi') or '未报告'}", ""]
    story = store.section(article_id, "story")
    if story:
        text.append("人工已核验" if story["human_verified"] else "人工待核验（原文匹配不代表结论正确）")
        for key, title in STORY_FIELDS.items():
            text.extend(["", "## " + title])
            for claim in story["content"].get(key, []):
                label = {"author_claim": "作者主张", "direct_result": "直接结果", "ai_inference": "推测/建议"}[claim["kind"]]
                text.append(f"- [{label}] {claim['text']} （证据：{', '.join(claim['evidence_ids']) or '无'}）")
        text.extend(["", "## 原文证据"])
        for evidence in story["evidence"]:
            text.append(f"- {evidence['evidence_id']} | {evidence['file_id']} | {evidence['loc_id']} | {evidence['figure'] or '图表未报告'}\n  > {evidence['quote']}")
    text.extend(["", "## 人工笔记", a["notes"] or "未记录"])
    text.extend(['', '手动分类：'+('；'.join(c['name'] for c in store.article_categories(article_id)) or '未分类')])
    text.extend(["", "## 原文高亮与批注"])
    for file in store.files(article_id):
        for mark in store.highlights(file["file_id"]):
            anchor = mark["anchors"][0]
            location = f"实际第{anchor['page_index']+1}页" if mark["kind"] == "pdf" else anchor["loc_id"]
            text.append(f"- {file['name']} | {location} | {mark['color']} | {mark['mark_id']}\n  > {mark['text']}\n  批注：{mark['note'] or '未记录'}\n  页内批注样式：{dumps(mark.get('appearance',{}))}")
    return "\n".join(text)

def excel_export(store, ids, target, cancel=lambda: False):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    book = Workbook()
    book.remove(book.active)
    headers = {
        "文献": ["article_id", "标题", "中文译名", "期刊", "年份", "DOI", "作者", "体系", "结构", "标签", "状态", "元数据JSON", "元数据候选来源", "人工笔记"],
        "附件": ["article_id", "file_id", "SHA256", "文件名", "角色", "相对路径", "解析状态", "说明", "配对依据"],
        "样品工艺": ["article_id", "sample_id", "样品标签", "组别", "步骤顺序", "步骤类别", "步骤描述", "参数", "原文值", "原文单位", "标准值", "标准单位", "计算基准", "换算说明", "冲突组", "证据ID", "人工核验"],
        "性能": ["article_id", "sample_id", "record_id", "测量类型", "条件JSON", "参数", "原文值", "原文单位", "标准值", "标准单位", "计算基准", "换算说明", "冲突组", "证据ID", "人工核验"],
        "证据": ["article_id", "分析段", "evidence_id", "file_id", "位置ID", "图表", "印刷页码", "原文片段", "引用匹配", "人工核验"],
        "阅读卡": ["article_id", "项目", "类别", "内容", "证据ID", "人工核验"],
        "修订历史": ["history_id", "article_id", "范围", "旧值JSON", "新值JSON", "来源", "UTC时间"],
        "高亮批注": ["article_id", "file_id", "mark_id", "文件名", "文件SHA256", "类型", "原文", "位置JSON", "颜色", "批注", "创建UTC", "修改UTC", "页内批注样式JSON"],
        "手动分类": ["article_id", "category_id", "类别", "归类UTC"],
    }
    sheets = {k: book.create_sheet(k) for k in headers}
    for name, cols in headers.items():
        sheets[name].append(cols)
    def safe(v):
        if isinstance(v, (dict, list)):
            v = dumps(v)
        if isinstance(v, str) and v.startswith(("=", "+", "-", "@")):
            v = "'" + v
        return v
    def append(name, row):
        # Excel cells have a 32767-character limit. Continue explicitly, never truncate silently.
        lengths = [len(v) if isinstance(v, str) else 0 for v in row]
        count = max(1, (max(lengths, default=0) + 29999) // 30000)
        for part in range(count):
            chunk = [v[part*30000:(part+1)*30000] if isinstance(v, str) and len(v)>30000 else v for v in row]
            sheets[name].append([safe(v) for v in chunk])
    def param(p):
        return [p[k] for k in ["name", "original_value", "original_unit", "normalized_value", "normalized_unit", "basis", "conversion_note", "conflict_group"]] + [", ".join(p["evidence_ids"])]
    for aid in ids:
        if cancel():
            raise InterruptedError("导出已取消")
        a = store.article(aid)
        m = a["metadata"]
        for category in store.rows('SELECT ac.*,c.name FROM article_categories ac JOIN categories c USING(category_id) WHERE article_id=?',(aid,)):
            append('手动分类',[aid,category['category_id'],category['name'],category['assigned']])
        append("文献", [aid, m.get("title"), m.get("title_zh"), m.get("journal"), m.get("year"), m.get("doi"), "; ".join(m.get("authors", [])), m.get("material"), m.get("structure"), "; ".join(m.get("tags", [])), next((r["state"] for r in store.list_articles() if r["article_id"]==aid), ""), dumps(m), dumps(a["candidates"]), a["notes"]])
        for f in store.files(aid):
            append("附件", [aid] + [f[k] for k in ["file_id", "sha256", "name", "role", "path", "status", "reason", "pairing"]])
            for mark in store.highlights(f["file_id"]):
                append("高亮批注", [aid,f["file_id"],mark["mark_id"],f["name"],mark["file_sha256"],mark["kind"],mark["text"],dumps(mark["anchors"]),mark["color"],mark["note"],mark["created"],mark["modified"],dumps(mark.get('appearance',{}))])
        for name in ["story", "process"]:
            section = store.section(aid, name)
            if not section:
                continue
            verified = "已人工核验" if section["human_verified"] else "待人工核验"
            checks = {c["evidence_id"]: c["matched"] for c in section["checks"]}
            for e in section["evidence"]:
                append("证据", [aid, name] + [e[k] for k in ["evidence_id", "file_id", "loc_id", "figure", "printed_page", "quote"]] + ["匹配" if checks.get(e["evidence_id"]) else "不匹配", verified])
            if name == "story":
                for k, claims in section["content"].items():
                    for claim in claims:
                        append("阅读卡", [aid, STORY_FIELDS[k], claim["kind"], claim["text"], ", ".join(claim["evidence_ids"]), verified])
            else:
                for sample in section["content"]["samples"]:
                    if not sample["steps"]:
                        append("样品工艺", [aid, sample["sample_id"], sample["label"], sample["group"]] + [None]*12 + [verified])
                    for step in sample["steps"]:
                        base = [aid, sample["sample_id"], sample["label"], sample["group"], step["order"], step["category"], step["description"]]
                        for p in step["parameters"]:
                            append("样品工艺", base + param(p) + [verified])
                        if not step["parameters"]:
                            append("样品工艺", base + [None]*8 + [", ".join(step["evidence_ids"]), verified])
                for rec in section["content"]["performance"]:
                    for p in rec["values"]:
                        append("性能", [aid, rec["sample_id"], rec["record_id"], rec["measurement_kind"], dumps(rec["conditions"])] + param(p) + [verified])
        for h in store.rows("SELECT * FROM history WHERE article_id=? ORDER BY created", (aid,)):
            append("修订历史", [h[k] for k in ["history_id", "article_id", "scope", "old_value", "new_value", "source", "created"]])
    for sheet in book:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="214D62")
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = min(60, max(16, len(str(col[0].value)) * 2 + 6))
    target = Path(target)
    temporary = target.with_name(target.name + ".writing")
    try:
        book.save(temporary)
        temporary.replace(target)
    finally:
        temporary.unlink(missing_ok=True)
